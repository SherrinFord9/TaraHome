#!/usr/bin/env python3
import importlib.util
import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

import openpyxl

spec = importlib.util.spec_from_file_location('search_change', Path(__file__).with_name('analyze-search-change.py'))
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def add_table(book, name, rows):
    sheet = book.create_sheet(name)
    sheet.append([name] + [f'{period} {metric}' for metric in ('Clicks', 'Impressions', 'CTR', 'Position')
                          for period in ('7/20/26 - 7/26/26', '7/12/26 - 7/18/26')])
    for row in rows:
        sheet.append([*row, 0, 0, 0, 0])
    return sheet


class SearchChangeTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.repo = Path(self.directory.name)
        module.git(self.repo, 'init', '-q')
        module.git(self.repo, 'config', 'user.name', 'Fixture')
        module.git(self.repo, 'config', 'user.email', 'fixture@example.com')
        for slug in ('a', 'b', 'c', 'removed'):
            file = self.repo / 'blog' / slug / 'index.html'
            file.parent.mkdir(parents=True)
            file.write_text(f'<title>{slug}</title><p>Original</p>')
        (self.repo / 'blog/index.html').write_text('<title>Library</title>')
        self.commit()
        (self.repo / 'blog/a/index.html').write_text('<title>New title</title><p>Original</p>')
        (self.repo / 'blog/b/index.html').write_text('<title>b</title><p>Revised body</p>')
        (self.repo / 'blog/removed/index.html').unlink()
        (self.repo / 'blog/new').mkdir()
        (self.repo / 'blog/new/index.html').write_text('<title>New page</title>')
        self.commit()
        self.ref = module.git(self.repo, 'rev-parse', 'HEAD').strip()
        self.book = openpyxl.Workbook()
        self.book.remove(self.book.active)
        add_table(self.book, 'Pages', [
            ['https://tarahome.ai/blog/a/', 1, 1, 4, 20],
            ['https://tarahome.ai/blog/b/', 0, 1, 3, 30],
            ['https://tarahome.ai/blog/c/', 0, 2, 5, 50],
            ['https://tarahome.ai/blog/removed/', 0, 1, 0, 5],
            ['https://tarahome.ai/blog/new/', 0, 0, 2, 0],
            ['https://tarahome.ai/blog/', 0, 0, 1, 5],
        ])
        add_table(self.book, 'Devices', [['Desktop', 1, 5, 12, 108]])
        add_table(self.book, 'Countries', [['United States', 1, 5, 12, 108]])
        filters = self.book.create_sheet('Filters')
        filters.append(['Filter', 'Value'])
        filters.append(['Search type', 'Web'])
        filters.append(['Date', 'Jul 20, 2026-Jul 26, 2026'])
        self.source = self.repo / 'comparison.xlsx'
        self.performance = self.repo / 'performance.json'
        daily = [{'date': (date(2026, 7, 12) + timedelta(days=n)).isoformat(),
                  'clicks': 5 if n == 0 else 1 if n == 8 else 0,
                  'impressions': 108 if n == 0 else 12 if n == 8 else 0} for n in range(15)]
        self.performance.write_text(json.dumps({'daily': daily}))

    def commit(self):
        module.git(self.repo, 'add', 'blog')
        module.git(self.repo, 'commit', '-qm', 'fixture')

    def analyze(self):
        self.book.save(self.source)
        return module.analyze(self.source, self.performance, self.repo, self.ref)

    def test_cohorts_use_history_not_worktree_and_keep_totals_separate(self):
        (self.repo / 'blog/c/index.html').write_text('<title>Uncommitted later edit</title>')
        result = self.analyze()
        cohorts = result['articleCohortsWithBaselineImpressions']
        self.assertEqual(cohorts['titleChanged']['before']['impressions'], 20)
        self.assertEqual(cohorts['otherHtmlChanged']['before']['impressions'], 30)
        self.assertEqual(cohorts['unchangedHtml']['before']['impressions'], 50)
        self.assertEqual(cohorts['unchangedHtml']['rows'], 1)
        self.assertEqual(cohorts['notPresentInBothRevisions']['rows'], 1)
        self.assertEqual(result['siteTotalsFromDailySeries']['after']['impressions'], 12)
        self.assertEqual(result['pageAggregationTotalsNotSiteTotals']['after']['impressions'], 15)
        self.assertEqual(len(result['pages']), 6)

    def test_cumulative_and_inconsistent_periods_are_rejected(self):
        self.book['Pages']['B1'] = 'Clicks'
        with self.assertRaisesRegex(ValueError, 'dated comparison'):
            self.analyze()
        self.book['Pages']['B1'] = '7/20/26 - 7/27/26 Clicks'
        with self.assertRaisesRegex(ValueError, 'two non-overlapping'):
            self.analyze()

    def test_filtered_or_unreconciled_exports_are_rejected(self):
        self.book['Devices']['D2'] = 13
        with self.assertRaisesRegex(ValueError, 'Segment totals'):
            self.analyze()
        self.book['Devices']['D2'] = 12
        self.book['Filters'].append(['Country', 'United States'])
        with self.assertRaisesRegex(ValueError, 'unfiltered Web'):
            self.analyze()

    def test_missing_daily_dates_are_not_zeroes(self):
        data = json.loads(self.performance.read_text())
        data['daily'].pop(3)
        self.performance.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'does not cover'):
            self.analyze()

    def test_duplicate_rows_and_bad_counts_are_rejected(self):
        self.book['Pages']['D2'] = -1
        with self.assertRaisesRegex(ValueError, 'Invalid comparison count'):
            self.analyze()
        self.book['Pages']['D2'] = 4
        self.book['Pages']['A3'] = self.book['Pages']['A2'].value
        with self.assertRaisesRegex(ValueError, 'duplicate segment'):
            self.analyze()

    def test_title_html_entities_and_whitespace_are_normalized(self):
        self.assertEqual(module.title('<title>A &amp;\n B</title>'), 'A & B')

    def test_fragment_rows_are_preserved_without_merging(self):
        self.book['Pages'].append(['https://tarahome.ai/#camera-kit', 0, 0, 1, 2, 0, 0, 0, 0])
        result = self.analyze()
        self.assertEqual(result['pages'][-1]['cohort'], 'nonCanonicalPageVariant')
        self.assertEqual(result['pages'][-1]['label'], 'https://tarahome.ai/#camera-kit')
        self.assertEqual(result['siteTotalsFromDailySeries']['after']['impressions'], 12)

    def test_unequal_and_cross_sheet_periods_are_rejected(self):
        sheet = self.book['Pages']
        for cell in sheet[1][1:]:
            cell.value = cell.value.replace('7/26/26', '7/27/26')
        with self.assertRaisesRegex(ValueError, 'equal positive lengths'):
            self.analyze()
        for cell in sheet[1][1:]:
            cell.value = cell.value.replace('7/20/26', '7/21/26')
        with self.assertRaisesRegex(ValueError, 'Periods differ across sheets'):
            self.analyze()


if __name__ == '__main__':
    unittest.main()
